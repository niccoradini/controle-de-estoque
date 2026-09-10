#!/usr/bin/env python3
"""Gera a migração de séries das planilhas das lojas da rede."""

from datetime import date, datetime
from pathlib import Path
import sys

from openpyxl import load_workbook


STORE_CODES = {
    "89MN": "bq-lucas",
    "283H": "patio",
    "210H": "avenida",
}


def sql_text(value):
    return "'" + str(value or "").replace("'", "''") + "'"


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value or "")[:10]


def stock_status(deposit, system_status):
    deposit = str(deposit or "").strip().upper()
    system_status = str(system_status or "").strip().upper()
    if deposit == "RPAR":
        return "repair"
    if system_status == "DEPS NREM" or deposit in {"DEPS", "NREM"}:
        return "incoming"
    if system_status == "DEPS" and deposit in {"EXPO", "LOJA", "LVUT"}:
        return "available"
    return "ignored"


def main():
    if len(sys.argv) < 3:
        raise SystemExit("uso: build-network-serials.py saida.sql planilha...")
    output = Path(sys.argv[1])
    rows = []
    for source_name in sys.argv[2:]:
        source = Path(source_name)
        sheet = load_workbook(source, read_only=True, data_only=True).active
        headers = {str(cell.value or "").strip(): index for index, cell in enumerate(next(sheet.iter_rows()), start=1)}
        for values in sheet.iter_rows(min_row=2, values_only=True):
            def field(name):
                return values[headers[name] - 1]
            material = str(field("Material") or "").strip()
            serial = str(field("Nº de série") or "").strip()
            center = str(field("Centro") or "").strip().upper()
            if not material or not serial or center not in STORE_CODES:
                continue
            rows.append((
                STORE_CODES[center], material, serial,
                stock_status(field("Depósito"), field("Status sistema")),
                iso_date(field("Modificado em")), source.name,
            ))

    insert_header = """INSERT INTO network_inventory_serials
  (store_code, material_code, serial_number, stock_status, modified_on, source_file)
VALUES
"""
    inserts = []
    for start in range(0, len(rows), 200):
        values = ",\n".join(
            "  (" + ", ".join(sql_text(value) for value in row) + ")"
            for row in rows[start:start + 200]
        )
        inserts.append(insert_header + values + ";")

    sql = f"""PRAGMA foreign_keys = ON;

CREATE TABLE IF NOT EXISTS network_inventory_serials (
  store_code TEXT NOT NULL REFERENCES network_stores(code) ON DELETE CASCADE,
  material_code TEXT NOT NULL COLLATE NOCASE,
  serial_number TEXT NOT NULL COLLATE NOCASE,
  stock_status TEXT NOT NULL CHECK (stock_status IN ('available', 'incoming', 'repair', 'ignored')),
  modified_on TEXT NOT NULL,
  source_file TEXT NOT NULL,
  PRIMARY KEY (store_code, serial_number)
);

CREATE INDEX IF NOT EXISTS idx_network_inventory_serials_material
ON network_inventory_serials (store_code, material_code, stock_status);

DELETE FROM network_inventory_serials
WHERE store_code IN ('bq-lucas', 'patio', 'avenida');

{chr(10).join(inserts)}
"""
    output.write_text(sql, encoding="utf-8")
    print({"rows": len(rows), "output": str(output)})


if __name__ == "__main__":
    main()
