from pathlib import Path
from openpyxl import load_workbook

SOURCE = Path('/workspace/scratch/a92edb98aecd/promo-materials/new_9307_portfolio-outlet-set26-semsmartsxlsxvnd-1789765651.xlsx')
OUTPUT = Path(__file__).resolve().parents[1] / 'migrations/0097_promotions_page_2026_09_22.sql'


def sql(value):
    if value is None:
        return 'NULL'
    if isinstance(value, (int, float)):
        return str(int(value))
    return "'" + str(value).replace("'", "''") + "'"


def cluster(category, line, name):
    text = f'{category} {line} {name}'.lower()
    if 'smartphone' in text:
        return 'devices'
    if 'wearable' in text or 'watch' in text:
        return 'wearables'
    if 'carregador' in text or 'powerbank' in text:
        return 'chargers'
    if 'capa' in text or 'case' in text:
        return 'cases'
    if 'película' in text or 'filme' in text:
        return 'screen_protectors'
    if 'áudio' in text or 'caixa de som' in text or 'fone' in text:
        return 'speakers'
    return 'misc'


wb = load_workbook(SOURCE, read_only=True, data_only=True)
rows = []
for position, row in enumerate(wb['Portfólio Elegível 16.09'].iter_rows(min_row=2, values_only=True), 1):
    material, name, category, line, _, _, discount = row[:7]
    if material is None or not name:
        continue
    discount_percent = round(float(discount) * 100) if discount is not None else None
    rows.append({
        'material': str(material).strip(),
        'name': str(name).strip(),
        'category': cluster(category, line, name),
        'group': str(category or 'Produto').strip(),
        'discount': discount_percent,
        'discount_text': f'{discount_percent}% OFF' if discount_percent is not None else 'Oferta',
        'regular_price': None,
        'price': None,
        'installments': None,
        'installment_price': None,
        'conditions': f'{discount_percent}% de desconto conforme o portfólio elegível de setembro de 2026.' if discount_percent is not None else None,
        'valid_from': None,
        'valid_to': None,
        'notes': None,
        'source': 'Portfólio Outlet setembro/2026 - base 16/09',
        'sort': 1000 + position,
    })


explicit = [
    ('22023045', 'Apple AirPods 4', 'misc', 'Apple', None, 'Preço especial', None, 79900, None, None, 'Elegível na compra com iPhone das famílias 14, 15, 16, 16e, 17, 17e e Air.', None, None, None, 'Campanha Apple enviada em 21/09', 10),
    ('22023046', 'Apple AirPods 4 com Cancelamento Ativo de Ruído', 'misc', 'Apple', None, 'Preço especial', None, 119900, None, None, 'Elegível na compra com iPhone das famílias 14, 15, 16, 16e, 17, 17e e Air.', None, None, None, 'Campanha Apple enviada em 21/09', 11),
    ('22024320', 'Apple AirPods Pro 3', 'misc', 'Apple', None, 'Preço especial', None, 169900, None, None, 'Elegível na compra com iPhone das famílias 14, 15, 16, 16e, 17, 17e e Air.', None, None, None, 'Campanha Apple enviada em 21/09', 12),
    ('22024315', 'Apple Watch SE 3 GPS 40mm', 'wearables', 'Apple', None, 'Preço especial', None, 169900, None, None, 'Elegível na compra com iPhone das famílias 14, 15, 16, 16e, 17, 17e e Air.', None, None, None, 'Campanha Apple enviada em 21/09', 13),
    ('22024316', 'Apple Watch SE 3 GPS 44mm', 'wearables', 'Apple', None, 'Preço especial', None, 189900, None, None, 'Elegível na compra com iPhone das famílias 14, 15, 16, 16e, 17, 17e e Air.', None, None, None, 'Campanha Apple enviada em 21/09', 14),
    ('22024877', 'Console PS5 Slim Digital', 'misc', 'Games', None, 'R$ 200 OFF', 459900, 439900, 12, 36700, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', 'Imagens meramente ilustrativas. Consulte a tabela vigente para alterações de preço.', 'Ofertas Mês do Cliente 15 a 28/09', 20),
    ('22024878', 'Console PS5 Slim Disk', 'misc', 'Games', None, 'R$ 200 OFF', 499900, 479900, 12, 40000, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', 'Imagens meramente ilustrativas. Consulte a tabela vigente para alterações de preço.', 'Ofertas Mês do Cliente 15 a 28/09', 21),
    ('22024185', 'Samsung Galaxy Tab S10 Lite Wi-Fi 128GB', 'misc', 'Tablet', None, 'R$ 400 OFF', 299900, 259900, 12, 21700, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 22),
    ('22024765', 'Samsung Galaxy Tab A11+ 5G 128GB', 'misc', 'Tablet', None, 'R$ 400 OFF', 199900, 159900, 12, 13400, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 23),
    ('22023877', 'Powerbank Pocket 5.000 USB-C i2GO', 'chargers', 'Essenciais', None, 'R$ 20 OFF', 13900, 11900, 3, 4000, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 24),
    ('22023495', 'Powerbank 5.000mAh MagSafe Slim Ovvi', 'chargers', 'Essenciais', None, 'R$ 40 OFF', 39900, 35900, 3, 12000, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 25),
    ('22021678', 'Carregador Parede Universal GAN 65W Ovvi', 'chargers', 'Essenciais', None, 'R$ 30 OFF', 29900, 26900, 3, 9000, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 26),
    ('22022669', 'Carregador Parede Universal GAN 65W Ovvi', 'chargers', 'Essenciais', None, 'R$ 30 OFF', 29900, 26900, 3, 9000, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 27),
    ('22020879', 'Aspirador Robô Smart Positivo', 'misc', 'Casa Inteligente', None, 'R$ 200 OFF', 79900, 59900, 6, 10000, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 28),
    ('22022496', 'Samsung The Freestyle 2ª Geração', 'misc', 'Eletrônicos', None, 'R$ 1.200 OFF', 369900, 249900, 12, 20900, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 29),
    ('22022605', 'Repetidor Vivo Wi-Fi 6 com base removível - preto', 'misc', 'Conectividade', None, 'R$ 50 OFF', 44900, 39900, 3, 13300, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 30),
    ('22023416', 'Repetidor Vivo Wi-Fi 6 com base removível - branco', 'misc', 'Conectividade', None, 'R$ 50 OFF', 44900, 39900, 3, 13300, 'Preço promocional à vista. Parcelamento conforme as regras vigentes no Conteúdos Vivo.', '2026-09-15', '2026-09-28', None, 'Ofertas Mês do Cliente 15 a 28/09', 31),
    ('22024486', 'Nintendo Switch OLED com jogo', 'misc', 'Games', None, 'Até R$ 800 OFF', None, None, None, None, 'R$ 200 OFF no Vivo Total Essencial e Pro; R$ 400 OFF no Vivo Total Ultra, Família 2 e Família 3; R$ 800 OFF no Vivo Total Família 4, Família 5, V e Segmento V.', None, None, 'Condição aplicável aos modelos Nintendo Switch elegíveis informados na campanha.', 'Condição Nintendo enviada em 21/09', 32),
    ('22023954', 'Nintendo Switch OLED', 'misc', 'Games', None, 'Até R$ 800 OFF', None, None, None, None, 'R$ 200 OFF no Vivo Total Essencial e Pro; R$ 400 OFF no Vivo Total Ultra, Família 2 e Família 3; R$ 800 OFF no Vivo Total Família 4, Família 5, V e Segmento V.', None, None, 'Condição aplicável aos modelos Nintendo Switch elegíveis informados na campanha.', 'Condição Nintendo enviada em 21/09', 33),
]

by_material = {row['material']: row for row in rows}
for values in explicit:
    keys = ('material', 'name', 'category', 'group', 'discount', 'discount_text', 'regular_price', 'price', 'installments', 'installment_price', 'conditions', 'valid_from', 'valid_to', 'notes', 'source', 'sort')
    by_material[values[0]] = dict(zip(keys, values))
rows = list(by_material.values())

columns = ('material', 'name', 'category', 'group', 'discount', 'discount_text', 'regular_price', 'price', 'installments', 'installment_price', 'conditions', 'valid_from', 'valid_to', 'notes', 'source', 'sort')
values_sql = ',\n'.join('  (' + ', '.join(sql(row[key]) for key in columns) + ')' for row in rows)

OUTPUT.write_text(f'''PRAGMA foreign_keys = ON;

CREATE TABLE IF NOT EXISTS promotion_rules (
  material_code TEXT PRIMARY KEY,
  product_name TEXT NOT NULL,
  category TEXT NOT NULL,
  promotion_group TEXT NOT NULL,
  discount_percent INTEGER,
  discount_text TEXT NOT NULL,
  regular_price_cents INTEGER,
  promotional_price_cents INTEGER,
  installment_count INTEGER,
  installment_price_cents INTEGER,
  conditions TEXT,
  valid_from TEXT,
  valid_to TEXT,
  notes TEXT,
  source_name TEXT NOT NULL,
  sort_order INTEGER NOT NULL DEFAULT 1000,
  active INTEGER NOT NULL DEFAULT 1 CHECK (active IN (0, 1))
);

DELETE FROM promotion_rules;

INSERT INTO promotion_rules
  (material_code, product_name, category, promotion_group, discount_percent,
   discount_text, regular_price_cents, promotional_price_cents, installment_count,
   installment_price_cents, conditions, valid_from, valid_to, notes, source_name, sort_order)
VALUES
{values_sql};
''', encoding='utf-8')
print(f'Generated {OUTPUT} with {len(rows)} promotion rules')
