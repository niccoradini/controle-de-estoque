from datetime import datetime
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
UPLOAD = ROOT.parent / "upload"
OUTPUT = ROOT / "migrations" / "0092_renova_values_2026_09.sql"


def sql(value):
    return "'" + str(value).replace("'", "''") + "'"


def iso_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value)


def match_key(value):
    normalized = unicodedata.normalize("NFKD", str(value))
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = re.sub(r"\b5G\b", "", normalized.upper())
    normalized = normalized.replace("+", " PLUS ")
    return re.sub(r"[^A-Z0-9]+", "", normalized)


voucher_path = next(UPLOAD.glob("*voucher*.xlsx"))
voucher_sheet = load_workbook(voucher_path, data_only=True, read_only=True)["SAP"]
voucher_rows = []
for row in voucher_sheet.iter_rows(min_row=3, values_only=True):
    _, product_type, manufacturer, model, good, defective, table_date = row[:7]
    if model is None:
        continue
    source_key = re.sub(r"\s+", " ", str(model).strip()).upper()
    if not source_key.startswith("ASSURANT "):
        continue
    voucher_rows.append(
        (
            source_key,
            source_key.removeprefix("ASSURANT "),
            str(manufacturer).strip().upper(),
            str(product_type).strip().upper(),
            round(float(good) * 100),
            round(float(defective) * 100),
            iso_date(table_date),
            1,
        )
    )

boost_path = next(UPLOAD.glob("*boost*.xlsx"))
boost_sheet = load_workbook(boost_path, data_only=True, read_only=True)["Planilha1"]
boost_rows = []
for row in boost_sheet.iter_rows(min_row=2, values_only=True):
    manufacturer, device_name, bonus, channel, starts_on, ends_on, _ = row[:7]
    if device_name is None:
        continue
    if str(channel).strip().casefold() != "físico":
        continue
    boost_rows.append(
        (
            str(manufacturer).strip(),
            str(device_name).strip(),
            match_key(device_name),
            round(float(bonus) * 100),
            iso_date(starts_on),
            iso_date(ends_on),
            "Tabela Boost Vivo Renova setembro/2026 v5",
            1,
        )
    )

assert len(voucher_rows) == 916
assert len({row[0].casefold() for row in voucher_rows}) == len(voucher_rows)
assert {row[3] for row in voucher_rows} == {"SMARTPHONE", "TABLET"}
assert {row[6] for row in voucher_rows} == {"2026-09-15"}
assert len(boost_rows) == 85
assert len({row[2] for row in boost_rows}) == len(boost_rows)


def values(rows):
    return ",\n".join("  (" + ", ".join(sql(v) if isinstance(v, str) else str(v) for v in row) + ")" for row in rows)


def batched_inserts(table, columns, rows, batch_size=100):
    statements = []
    for start in range(0, len(rows), batch_size):
        batch = rows[start : start + batch_size]
        statements.append(
            f"INSERT INTO {table}\n  ({columns})\nVALUES\n{values(batch)};"
        )
    return "\n\n".join(statements)


content = f"""-- Valores oficiais do Vivo Renova recebidos em setembro de 2026.
-- Voucher: parceiro ASSURANT, tabela SAP de 15/09/2026.
-- Boost: tabela setembro/2026 v5, com 85 modelos do canal Físico.

DROP TABLE IF EXISTS renova_trade_in_values_20260915;
CREATE TABLE renova_trade_in_values_20260915 (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  source_key TEXT NOT NULL COLLATE NOCASE UNIQUE,
  device_name TEXT NOT NULL COLLATE NOCASE,
  manufacturer TEXT NOT NULL COLLATE NOCASE,
  product_type TEXT NOT NULL COLLATE NOCASE CHECK (product_type IN ('SMARTPHONE', 'TABLET')),
  good_cents INTEGER NOT NULL CHECK (good_cents >= 0),
  defective_cents INTEGER NOT NULL CHECK (defective_cents >= 0),
  table_date TEXT NOT NULL,
  active INTEGER NOT NULL DEFAULT 1 CHECK (active IN (0, 1))
);

{batched_inserts('renova_trade_in_values_20260915', 'source_key, device_name, manufacturer, product_type, good_cents, defective_cents, table_date, active', voucher_rows)}

ALTER TABLE renova_trade_in_values RENAME TO renova_trade_in_values_archive_20260804;
ALTER TABLE renova_trade_in_values_20260915 RENAME TO renova_trade_in_values;

DELETE FROM renova_manufacturer_boosts;
INSERT INTO renova_manufacturer_boosts
  (manufacturer, device_name, match_key, bonus_cents, starts_on, ends_on, source, active)
VALUES
{values(boost_rows)};

INSERT OR REPLACE INTO system_state (key, value) VALUES
  ('renova_voucher_table_date', '2026-09-15'),
  ('renova_voucher_source', 'Tabela Voucher Renova setembro/2026 v2'),
  ('renova_boost_table_date', '2026-09-17'),
  ('renova_boost_source', 'Tabela Boost Vivo Renova setembro/2026 v5');

CREATE TABLE _migration_0092_guard (valid INTEGER NOT NULL CHECK (valid = 1));
INSERT INTO _migration_0092_guard (valid)
SELECT CASE WHEN
  (SELECT COUNT(*) FROM renova_trade_in_values WHERE active = 1) = 916
  AND (SELECT COUNT(DISTINCT source_key) FROM renova_trade_in_values) = 916
  AND (SELECT COUNT(*) FROM renova_trade_in_values WHERE table_date = '2026-09-15') = 916
  AND (SELECT COUNT(*) FROM renova_manufacturer_boosts WHERE active = 1) = 85
  AND (SELECT COUNT(*) FROM renova_manufacturer_boosts WHERE ends_on >= '2026-09-21') = 85
  AND (SELECT bonus_cents FROM renova_manufacturer_boosts WHERE match_key = 'IPHONE18PROMAX2TB') = 80000
  AND (SELECT bonus_cents FROM renova_manufacturer_boosts WHERE match_key = 'SAMSUNGGALAXYS26FE256GB') = 37000
  THEN 1 ELSE 0 END;
DROP TABLE _migration_0092_guard;
"""

OUTPUT.write_text(content, encoding="utf-8")
print(f"Wrote {OUTPUT} with {len(voucher_rows)} vouchers and {len(boost_rows)} boosts")
